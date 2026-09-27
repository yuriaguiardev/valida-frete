"""Gera a apresentação (docs/entrega/apresentacao.pptx) a partir do código.

Uso (na raiz do projeto):
    pip install -r requirements-docs.txt
    python scripts/gerar_slides.py

Requer os diagramas já gerados (python scripts/gerar_afn.py) e o Graphviz
para os mini-diagramas da construção de Thompson.
"""

import subprocess
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from validafrete.afn import construir_afn  # noqa: E402
from validafrete.casos_teste import CASOS  # noqa: E402
from validafrete.expressoes import EXPRESSOES  # noqa: E402

# Paleta: azul-marinho de frota (dominante), âmbar de sinalização (destaque).
MARINHO = "1B2A41"
MARINHO_2 = "2C3E5C"
AMBAR = "F2A541"
BRANCO = "FFFFFF"
TINTA = "EEF2F7"
CINZA = "5B6B7F"
TEXTO = "1F2933"
VERDE = "1E7B45"
VERDE_CLARO = "E3F4E8"
VERMELHO = "B42318"
VERMELHO_CLARO = "FDECEA"
CODIGO_FUNDO = "0F1E2E"
CODIGO_TEXTO = "E3F0FF"

TITULO = "Cambria"
CORPO = "Calibri"
MONO = "Courier New"

IMAGENS = RAIZ / "docs" / "imagens"
AFN_PNG = RAIZ / "docs" / "afn"
SAIDA = RAIZ / "docs" / "entrega" / "apresentacao.pptx"
TMP = Path(tempfile.mkdtemp(prefix="validafrete_slides_"))


# ---------------------------------------------------------------------------
# Auxiliares de desenho
# ---------------------------------------------------------------------------

def cor(hex_):
    return RGBColor.from_string(hex_)


def fundo(slide, hex_):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = cor(hex_)


def caixa(slide, x, y, w, h, preenchimento, arredondada=True, raio=0.08, borda=None):
    forma = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if arredondada else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if arredondada:
        forma.adjustments[0] = raio
    forma.fill.solid()
    forma.fill.fore_color.rgb = cor(preenchimento)
    if borda:
        forma.line.color.rgb = cor(borda)
        forma.line.width = Pt(1)
    else:
        forma.line.fill.background()
    forma.shadow.inherit = False
    return forma


def texto(slide, x, y, w, h, paragrafos, tamanho=16, cor_=TEXTO, fonte=CORPO, negrito=False,
          alinhamento=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP, margem=0.0, espaco_depois=0,
          marcadores=False, italico=False):
    """Caixa de texto. ``paragrafos`` é uma string, uma lista de strings ou uma
    lista de listas de trechos (texto, {opções}) para formatação mista."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    for lado in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, lado, Inches(margem))
    if isinstance(paragrafos, str):
        paragrafos = [paragrafos]
    for i, par in enumerate(paragrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinhamento
        if espaco_depois:
            p.space_after = Pt(espaco_depois)
        if marcadores:
            _marcador(p)
        trechos = [(par, {})] if isinstance(par, str) else par
        for conteudo, opcoes in trechos:
            r = p.add_run()
            r.text = conteudo
            f = r.font
            f.name = opcoes.get("fonte", fonte)
            f.size = Pt(opcoes.get("tamanho", tamanho))
            f.bold = opcoes.get("negrito", negrito)
            f.italic = opcoes.get("italico", italico)
            f.color.rgb = cor(opcoes.get("cor", cor_))
    return tb


def _marcador(paragrafo):
    pPr = paragrafo._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.22)))
    pPr.set("indent", str(-Inches(0.22)))
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


def titulo(slide, principal, sobre=None, claro=False):
    """Título do slide; ``sobre`` é uma etiqueta pequena acima (ex.: ER-01)."""
    y = 0.45
    if sobre:
        texto(slide, 0.6, 0.35, 8, 0.35, sobre, tamanho=13, negrito=True, cor_=AMBAR if claro else CINZA)
        y = 0.65
    texto(slide, 0.6, y, 12.1, 0.8, principal, tamanho=32, negrito=True, fonte=TITULO,
          cor_=BRANCO if claro else MARINHO)


def chip(slide, x, y, w, cadeia, aceita, tamanho=11, h=0.34, espaco_visivel=True):
    caixa(slide, x, y, w, h, VERDE_CLARO if aceita else VERMELHO_CLARO, raio=0.3)
    marca = "✓ " if aceita else "✗ "
    exibida = "ε (vazia)" if cadeia == "" else (cadeia.replace(" ", "␣") if espaco_visivel else cadeia)
    texto(slide, x + 0.1, y, w - 0.15, h, [[(marca, {"fonte": CORPO, "negrito": True,
                                                    "cor": VERDE if aceita else VERMELHO}),
                                           (exibida, {"fonte": MONO})]],
          tamanho=tamanho, ancora=MSO_ANCHOR.MIDDLE, cor_=TEXTO)


def imagem(slide, caminho, x, y, w, h, alinhar="centro", escala=None):
    """Insere a imagem ajustada à caixa (ou numa escala fixa), preservando a proporção."""
    with Image.open(caminho) as im:
        largura, altura = im.size
    if escala is None:
        escala = min(w / largura, h / altura)
    lw, lh = largura * escala, altura * escala
    px = x + (w - lw) / 2 if alinhar == "centro" else x
    py = y + (h - lh) / 2 if alinhar == "centro" else y
    return slide.shapes.add_picture(str(caminho), Inches(px), Inches(py), Inches(lw), Inches(lh))


def escala_comum(caminhos, w, h):
    """Maior escala (polegadas por pixel) em que todas as imagens cabem em w × h."""
    escalas = []
    for c in caminhos:
        with Image.open(c) as im:
            escalas.append(min(w / im.size[0], h / im.size[1]))
    return min(escalas)


def notas(slide, conteudo):
    slide.notes_slide.notes_text_frame.text = conteudo


def rodape(slide, conteudo, claro=False):
    texto(slide, 0.6, 7.0, 12.1, 0.3, conteudo, tamanho=10, cor_="9FB0C6" if claro else CINZA)


# ---------------------------------------------------------------------------
# Imagens auxiliares
# ---------------------------------------------------------------------------

def mini_thompson(padrao, nome):
    dot = construir_afn(padrao).para_dot(padrao)
    dot = dot.replace('fontsize=11, width=0.45', 'fontsize=14, width=0.5')
    arquivo = TMP / f"{nome}.dot"
    arquivo.write_text(dot, encoding="utf-8")
    png = TMP / f"{nome}.png"
    subprocess.run(["dot", "-Tpng", "-Gdpi=220", str(arquivo), "-o", str(png)], check=True)
    return png


def fatias_afn(er, partes):
    """Divide diagramas muito largos em faixas, cortando entre estados."""
    caminho = AFN_PNG / f"{er.codigo.lower()}-{er.chave}.png"
    if partes == 1:
        return [caminho]
    with Image.open(caminho) as im:
        im = im.convert("RGB")
        largura, altura = im.size
        cinza = im.convert("L")
        pixels = cinza.load()
        cortes = [0]
        for k in range(1, partes):
            alvo = largura * k // partes
            janela = range(alvo - largura // 12, alvo + largura // 12)
            # coluna com menos pixels escuros (entre dois estados)
            melhor = min(janela, key=lambda cx: sum(pixels[cx, cy] < 200 for cy in range(0, altura, 2)))
            cortes.append(melhor)
        cortes.append(largura)
        arquivos = []
        for i in range(partes):
            faixa = im.crop((cortes[i], 0, cortes[i + 1], altura))
            destino = TMP / f"{er.chave}_{i}.png"
            faixa.save(destino)
            arquivos.append(destino)
    return arquivos


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def slide_capa(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, MARINHO)
    texto(s, 0.8, 0.8, 7, 0.4, "LINGUAGENS FORMAIS E AUTÔMATOS · TRABALHO DO 1º BIMESTRE",
          tamanho=13, negrito=True, cor_=AMBAR)
    texto(s, 0.8, 1.9, 7.2, 1.2, "ValidaFrete", tamanho=60, negrito=True, fonte=TITULO, cor_=BRANCO)
    texto(s, 0.8, 3.1, 6.8, 1.0, "Validador de cadastros de fretes com Expressões Regulares e AFNε",
          tamanho=22, cor_="CADCFC")
    texto(s, 0.8, 5.0, 6.5, 1.2, ["Yuri Aguiar  ·  Pedro Paulo  ·  João Rath",
                                  "Turma CC6NA  ·  30/09/2026"], tamanho=16, cor_=BRANCO, espaco_depois=6)
    exemplos = [("12.ABC.345/01DE-35", True), ("BRA2E19", True), ("ABC-1D23", False),
                ("(23) 91234-5678", False), ("31/04/2026", False), ("R$ 1.250,00", True)]
    for i, (cadeia, ok) in enumerate(exemplos):
        chip(s, 8.6, 1.5 + i * 0.72, 3.9, cadeia, ok, tamanho=15, h=0.52, espaco_visivel=False)
    notas(s, "YURI: Apresentar a equipe e o tema. O ValidaFrete valida cadastros de fretes usando "
             "cinco Expressões Regulares e mostra, com o AFNε, onde cada valor falhou.")


def slide_problema(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, "Cadastros de fretes chegam cheios de erros")
    texto(s, 0.6, 1.55, 4.6, 3.6, [
        "Transportadoras recebem planilhas preenchidas à mão por filiais, parceiros e clientes.",
        "Erros de formato só aparecem quando a nota fiscal é recusada ou o motorista não é encontrado.",
        [("Desde julho de 2026 ", {"negrito": True}),
         ("o CNPJ também pode ter letras (IN RFB 2.229/2024): validações antigas rejeitam cadastros "
          "corretos.", {})],
    ], tamanho=16, espaco_depois=12)
    erros = [
        ("CNPJ", "12ABC345/01DE-35", "máscara pela metade"),
        ("Placa", "ABC-1D23", "Mercosul não tem hífen"),
        ("Telefone", "(23) 91234-5678", "DDD 23 não existe"),
        ("Data", "31/04/2026", "abril tem só 30 dias"),
        ("Valor", "R$ 1.23,45", "grupo de milhar com 2 dígitos"),
    ]
    for i, (campo, valor, motivo) in enumerate(erros):
        y = 1.55 + i * 1.0
        caixa(s, 5.6, y, 7.1, 0.84, TINTA)
        texto(s, 5.8, y, 1.3, 0.84, campo, tamanho=15, negrito=True, cor_=MARINHO, ancora=MSO_ANCHOR.MIDDLE)
        chip(s, 7.1, y + 0.22, 2.9, valor, False, tamanho=12, h=0.4, espaco_visivel=False)
        texto(s, 10.2, y, 2.4, 0.84, motivo, tamanho=13, cor_=CINZA, ancora=MSO_ANCHOR.MIDDLE)
    notas(s, "YURI: Contextualizar o problema. Cada linha da direita é um erro real que o programa "
             "detecta; todos estão no arquivo dados/fretes_exemplo.csv.")


def slide_solucao(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, "A solução: entrada, processamento e saída")
    colunas = [
        ("Entrada", ["Valor digitado, com o tipo escolhido ou identificado automaticamente",
                     "Planilha CSV: id;cnpj;placa;telefone;data;valor"]),
        ("Processamento", ["re.fullmatch com a ER de cada campo",
                           "Se rejeitado: simulação do AFNε para achar o símbolo que falhou"]),
        ("Saída", ["Aceita / rejeitada, com posição e motivo",
                   "Resumo por ER, tabela de erros e relatório CSV para download"]),
    ]
    for i, (nome, itens) in enumerate(colunas):
        x = 0.6 + i * 4.25
        caixa(s, x, 1.7, 3.75, 3.7, TINTA)
        caixa(s, x + 0.3, 1.95, 0.62, 0.62, MARINHO, raio=0.5)
        texto(s, x + 0.3, 1.95, 0.62, 0.62, str(i + 1), tamanho=20, negrito=True, cor_=AMBAR,
              alinhamento=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
        texto(s, x + 1.1, 1.95, 2.5, 0.62, nome, tamanho=22, negrito=True, fonte=TITULO, cor_=MARINHO,
              ancora=MSO_ANCHOR.MIDDLE)
        texto(s, x + 0.3, 2.85, 3.2, 2.4, itens, tamanho=15, espaco_depois=10, marcadores=True)
        if i < 2:
            seta = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 3.8), Inches(3.35),
                                      Inches(0.4), Inches(0.4))
            seta.fill.solid()
            seta.fill.fore_color.rgb = cor(AMBAR)
            seta.line.fill.background()
    tecnologias = ["Python 3", "Flask", "módulo re", "pytest", "Graphviz"]
    for i, t in enumerate(tecnologias):
        caixa(s, 0.6 + i * 2.5, 5.9, 2.2, 0.55, MARINHO, raio=0.3)
        texto(s, 0.6 + i * 2.5, 5.9, 2.2, 0.55, t, tamanho=14, negrito=True, cor_=BRANCO,
              alinhamento=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
    notas(s, "YURI: Explicar as três funções da aplicação: validar um valor, processar uma planilha e "
             "simular o AFNε. Tecnologias: Python, Flask para a interface, re para as ERs, pytest e Graphviz.")


def slide_fonte_unica(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, "Uma única fonte para a mesma linguagem")
    texto(s, 0.6, 1.45, 12, 0.5,
          "Regra central do trabalho: ER formal, slides, código, testes e AFNε devem representar a mesma linguagem.",
          tamanho=16, cor_=CINZA, italico=True)
    caixa(s, 4.9, 3.25, 3.55, 1.45, MARINHO)
    texto(s, 4.9, 3.3, 3.55, 1.35, [[("expressoes.py", {"fonte": MONO, "negrito": True, "tamanho": 18,
                                                          "cor": AMBAR})],
                                    "5 ERs: formal, padrão, Σ, linguagem, limitações"],
          tamanho=13, cor_=BRANCO, alinhamento=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
    consumidores = [
        (0.6, 2.2, "validador.py + app.py", "re.fullmatch valida cada campo"),
        (0.6, 4.6, "casos_teste.py + tests/", "80 casos, 223 testes automatizados"),
        (9.0, 2.2, "afn.py", "Thompson gera o AFNε a partir do padrão do código"),
        (9.0, 4.6, "scripts/", "diagramas, fichas, relatório e estes slides"),
    ]
    for x, y, nome, desc in consumidores:
        caixa(s, x, y, 3.75, 1.3, TINTA)
        texto(s, x + 0.25, y + 0.15, 3.3, 1.0, [[(nome, {"fonte": MONO, "negrito": True, "cor": MARINHO})],
                                                 desc], tamanho=14, espaco_depois=4)
        linha = s.shapes.add_connector(1, Inches(x + 3.75 if x < 4 else x), Inches(y + 0.65),
                                       Inches(4.9 if x < 4 else 8.45), Inches(3.97))
        linha.line.color.rgb = cor(AMBAR)
        linha.line.width = Pt(2.5)
    caixa(s, 0.6, 6.2, 12.1, 0.6, VERDE_CLARO)
    texto(s, 0.8, 6.2, 11.8, 0.6, "Testes comparam AFNε × re.fullmatch em 15.000 cadeias geradas: 0 divergências.",
          tamanho=15, negrito=True, cor_=VERDE, ancora=MSO_ANCHOR.MIDDLE)
    notas(s, "YURI: Nada foi escrito duas vezes. O validador, os testes, os diagramas e a documentação leem "
             "as ERs do mesmo módulo. O AFNε não é desenhado à mão: é gerado do padrão que o programa usa.")


def slide_thompson(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, "Da Expressão Regular ao AFNε: construção de Thompson")
    blocos = [
        ("a", "Símbolo", "a"),
        ("a|b", "União r | s", "a|b"),
        ("ab", "Concatenação rs", "ab"),
        ("a*", "Fecho de Kleene r*", "a*"),
        ("a?", "Opcional r? = (r | ε)", "a?"),
        ("a{3}", "Repetição r{3} = rrr", "a{3}"),
    ]
    minis = [mini_thompson(padrao, f"mini_{i}") for i, (padrao, _, _) in enumerate(blocos)]
    escala = escala_comum(minis, 3.55, 1.6)
    for i, (padrao, nome, rotulo) in enumerate(blocos):
        col, lin = i % 3, i // 3
        x, y = 0.6 + col * 4.1, 1.5 + lin * 2.4
        caixa(s, x, y, 3.85, 2.25, TINTA)
        texto(s, x + 0.2, y + 0.1, 3.5, 0.4, [[(nome, {"negrito": True, "cor": MARINHO}),
                                               ("   " + rotulo, {"fonte": MONO, "cor": CINZA})]], tamanho=14)
        imagem(s, minis[i], x + 0.15, y + 0.55, 3.55, 1.6, escala=escala)
    texto(s, 0.6, 6.4, 12.1, 0.6, [
        [("Setas tracejadas = movimentos ε.  ", {"negrito": True}),
         ("Simplificações que preservam a linguagem: a concatenação funde o final de r com o início de s, "
          "e r? usa um único ε para pular r.", {})]], tamanho=14, cor_=CINZA)
    notas(s, "YURI: Explicar as regras de Thompson. Cada operador vira um fragmento com um estado inicial e "
             "um final. Estes mini-diagramas foram gerados pelo mesmo código que gera os AFNε das cinco ERs.")


def slide_ficha(prs, er, orador):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, er.nome, sobre=er.codigo)
    # coluna esquerda: alfabeto e linguagem
    caixa(s, 0.6, 1.55, 4.7, 5.3, TINTA)
    texto(s, 0.85, 1.7, 4.3, 0.35, "ALFABETO (Σ)", tamanho=12, negrito=True, cor_=CINZA)
    alfabeto = [a for a in er.alfabeto]
    texto(s, 0.85, 2.05, 4.3, 1.9, alfabeto, tamanho=12.5 if len(" ".join(alfabeto)) < 260 else 11,
          espaco_depois=3)
    texto(s, 0.85, 3.95, 4.3, 0.35, "LINGUAGEM L", tamanho=12, negrito=True, cor_=CINZA)
    texto(s, 0.85, 4.3, 4.3, 2.5, er.linguagem, tamanho=12.5)
    # coluna direita: formal, código e equivalências
    texto(s, 5.6, 1.55, 7, 0.3, "ER FORMAL", tamanho=12, negrito=True, cor_=CINZA)
    caixa(s, 5.6, 1.85, 7.1, 0.95, "E6EEF7")
    texto(s, 5.75, 1.85, 6.8, 0.95, er.formal, tamanho=15 if len(er.formal) < 90 else 13,
          fonte="Cambria Math", ancora=MSO_ANCHOR.MIDDLE)
    texto(s, 5.6, 2.9, 7, 0.3, "SINTAXE NO CÓDIGO (PYTHON)", tamanho=12, negrito=True, cor_=CINZA)
    altura_codigo = 1.3 if len(er.padrao) > 120 else 0.8
    caixa(s, 5.6, 3.2, 7.1, altura_codigo, CODIGO_FUNDO)
    texto(s, 5.75, 3.2, 6.85, altura_codigo, f're.fullmatch(r"{er.padrao}", cadeia)',
          tamanho=10.5 if len(er.padrao) > 120 else 12, fonte=MONO, cor_=CODIGO_TEXTO, ancora=MSO_ANCHOR.MIDDLE)
    y = 3.2 + altura_codigo + 0.15
    texto(s, 5.6, y, 7, 0.3, "EQUIVALÊNCIA", tamanho=12, negrito=True, cor_=CINZA)
    linhas = [e for e in er.equivalencias if e[0] != "re.fullmatch"]
    espaco = 6.95 - (y + 0.35)
    maximo = int(espaco // 0.4)
    linhas = linhas[:maximo]
    for i, (trecho, significado) in enumerate(linhas):
        ly = y + 0.35 + i * 0.4
        texto(s, 5.6, ly, 2.3, 0.4, trecho, tamanho=11, fonte=MONO, cor_=MARINHO, negrito=True,
              ancora=MSO_ANCHOR.MIDDLE)
        texto(s, 7.95, ly, 4.75, 0.4, significado, tamanho=11, ancora=MSO_ANCHOR.MIDDLE)
    notas(s, f"{orador}: Ler a finalidade: {er.finalidade} Mostrar que cada atalho da sintaxe do código "
             f"corresponde a um operador formal (tabela de equivalência). re.fullmatch exige que a cadeia "
             f"inteira pertença à linguagem, como as âncoras ^ e $.")


def slide_afn(prs, er, orador, partes=1, com_testes=True):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    afn = construir_afn(er.padrao)
    titulo(s, "AFNε e testes" if com_testes else "AFNε correspondente", sobre=f"{er.codigo} · {er.nome}")
    texto(s, 0.6, 1.3, 12.1, 0.35,
          f"{afn.num_estados} estados · {len(afn.transicoes)} transições · {afn.total_vazias} movimentos ε · "
          f"inicial q{afn.inicial} · final q{afn.final} (círculo duplo)", tamanho=13, cor_=CINZA)
    area_h = 2.75 if com_testes else 5.45
    faixas = fatias_afn(er, partes)
    h_faixa = area_h / len(faixas)
    escala = escala_comum(faixas, 12.1, h_faixa - 0.05)
    for i, faixa in enumerate(faixas):
        alinhar = "centro" if len(faixas) == 1 else "esquerda"
        imagem(s, faixa, 0.6, 1.75 + i * h_faixa, 12.1, h_faixa - 0.05, alinhar=alinhar, escala=escala)
    if er.chave == "telefone":
        _rotulos_telefone(s)
    if com_testes:
        _testes(s, er, 4.65)
    notas(s, f"{orador}: Mostrar o estado inicial (seta 'início'), o final (círculo duplo) e os movimentos ε "
             f"(tracejados). Percorrer uma cadeia aceita e uma rejeitada. Para ampliar o diagrama, abrir "
             f"Expressões e AFNε > {er.codigo} > 'Abrir diagrama em tela cheia' na aplicação.")
    return s


def _rotulos_telefone(s):
    """Etiquetas nas áreas livres ao redor do AFNε do telefone."""
    for x, y, rotulo in ((0.6, 2.3, "'(' K ')'  DDD com parênteses  →"), (0.6, 6.35, "↑ '+' 55 (␣ | ε)  opcional"),
                         (5.7, 6.6, "← K  DDD sem parênteses"), (8.4, 3.9, "↓ celular 9D{4} | fixo FD{3}")):
        caixa(s, x, y, 2.9, 0.36, AMBAR, raio=0.3)
        texto(s, x, y, 2.9, 0.36, rotulo, tamanho=11, negrito=True, cor_=MARINHO,
              alinhamento=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)


def _testes(s, er, y, grande=False):
    casos = CASOS[er.chave]
    aceitas = [c for c in casos if c.aceita]
    rejeitadas = [c for c in casos if not c.aceita]
    texto(s, 0.6, y, 6, 0.3, f"ACEITAS ({len(aceitas)})", tamanho=12, negrito=True, cor_=VERDE)
    texto(s, 6.75, y, 6, 0.3, f"REJEITADAS ({len(rejeitadas)})", tamanho=12, negrito=True, cor_=VERMELHO)
    for grupo, x0 in ((aceitas, 0.6), (rejeitadas, 6.75)):
        for i, caso in enumerate(grupo):
            col, lin = i % 2, i // 2
            if grande:
                chip(s, x0 + col * 3.0, y + 0.4 + lin * 0.62, 2.9, caso.cadeia, caso.aceita, tamanho=12, h=0.5)
            else:
                chip(s, x0 + col * 3.0, y + 0.35 + lin * 0.43, 2.9, caso.cadeia, caso.aceita, tamanho=10.5, h=0.36)
    limites = sum(c.limite for c in casos)
    texto(s, 0.6, 7.0, 12.1, 0.3,
          f"{len(casos)}/{len(casos)} corretos em re.fullmatch e no AFNε · {limites} casos-limite · "
          f"lista completa com descrições em docs/EXPRESSOES.md",
          tamanho=11, cor_=CINZA)


def slide_testes_tel(prs, er, orador):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, "Testes e o porquê de o DDD aparecer duas vezes", sobre=f"{er.codigo} · {er.nome}")
    caixa(s, 0.6, 1.55, 12.1, 1.6, TINTA)
    texto(s, 0.85, 1.7, 11.6, 1.4, [
        [("Parênteses balanceados: ", {"negrito": True, "cor": MARINHO}),
         ("a ER é  '(' K ')' | K  — e não  '('? K ')'? , que aceitaria  (11 91234-5678.", {})],
        [("Uma ER não tem memória para lembrar que abriu um parêntese; por isso a subexpressão K dos 67 DDDs "
          "é repetida, e o AFNε tem dois blocos de DDD.", {})],
    ], tamanho=15, espaco_depois=6)
    _testes(s, er, 3.5, grande=True)
    notas(s, f"{orador}: Explicar a duplicação do DDD: é o custo de exigir parênteses balanceados sem "
             f"memória. Destacar os casos-limite: DDD 99 (maior existente) e número com um dígito a menos.")


def slide_entradas(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, "Entradas vazias ou inválidas geram mensagens claras")
    casos = [
        ("Valor vazio", "Entrada vazia (ε): a palavra vazia não pertence à linguagem desta ER."),
        ("Símbolo fora de Σ", "O símbolo ␣ (posição 19) não pertence ao alfabeto Σ. Verifique se há "
                              "espaços sobrando no início ou no fim."),
        ("Falha no AFNε", "Rejeitada na posição 5: depois de ler \"31/0\", o AFNε não tem transição com o símbolo '4'."),
        ("Cadeia incompleta", "Toda a entrada foi lida, mas o AFNε terminou em estados não finais."),
        ("Arquivo", "Nenhum arquivo, extensão diferente de .csv, vazio, binário, maior que 1 MB ou sem registros."),
        ("Estrutura do CSV", "Colunas ausentes no cabeçalho; linha com colunas a mais ou a menos é "
                             "sinalizada e o restante continua sendo processado."),
    ]
    for i, (nome, msg) in enumerate(casos):
        col, lin = i % 2, i // 2
        x, y = 0.6 + col * 6.15, 1.55 + lin * 1.75
        caixa(s, x, y, 5.95, 1.55, TINTA)
        texto(s, x + 0.25, y + 0.15, 5.5, 0.4, nome, tamanho=16, negrito=True, cor_=MARINHO)
        texto(s, x + 0.25, y + 0.58, 5.5, 0.9, msg, tamanho=13, cor_=TEXTO)
    texto(s, 0.6, 6.85, 12.1, 0.4, "Espaços extras não são removidos: \" 11222333000181\" é rejeitado, exatamente como nos testes — "
              "o programa reconhece a mesma linguagem da ER.", tamanho=13, cor_=CINZA, italico=True)
    notas(s, "JOÃO: O diagnóstico usa o AFNε: o programa lê a cadeia símbolo a símbolo e informa onde o "
             "conjunto de estados ficou vazio. Não fazemos strip para não mudar a linguagem.")


def slide_demo(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, "Demonstração")
    passos = [
        ("1", "Cadeia aceita", "BRA2E19 → ER-02"),
        ("2", "Cadeia rejeitada", "(23) 91234-5678 → DDD inexistente"),
        ("3", "Planilha de exemplo", "15 fretes: 5 válidos, 10 com erros"),
        ("4", "Simulação no AFNε", "estados ativos a cada símbolo"),
    ]
    for i, (n, nome, desc) in enumerate(passos):
        y = 1.6 + i * 1.3
        caixa(s, 0.6, y, 0.6, 0.6, MARINHO, raio=0.5)
        texto(s, 0.6, y, 0.6, 0.6, n, tamanho=18, negrito=True, cor_=AMBAR, alinhamento=PP_ALIGN.CENTER,
              ancora=MSO_ANCHOR.MIDDLE)
        texto(s, 1.4, y - 0.05, 3.3, 0.4, nome, tamanho=16, negrito=True, cor_=MARINHO)
        texto(s, 1.4, y + 0.33, 3.3, 0.6, desc, tamanho=13, cor_=CINZA)
    imagem(s, IMAGENS / "tela_validar.png", 5.0, 1.45, 7.7, 2.75)  # recortada em docs/imagens
    imagem(s, IMAGENS / "tela_simulacao.png", 5.0, 4.3, 7.7, 2.6)
    rodape(s, "Executar: python -m validafrete  →  http://127.0.0.1:5050")
    notas(s, "YURI: Demonstração ao vivo. (1) Validar BRA2E19 em 'Identificar automaticamente'. "
             "(2) Validar (23) 91234-5678 e ler a mensagem. (3) Processar planilha > Usar arquivo de exemplo; "
             "baixar o relatório. (4) Expressões e AFNε > ER-02 > simular BRA-2E19 e mostrar o conjunto vazio no "
             "passo 6. Se o professor pedir outra entrada, digitar na página inicial.")


def slide_numeros(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, MARINHO)
    titulo(s, "Testes e análise dos resultados", claro=True)
    numeros = [("223", "testes automatizados\n(pytest)"), ("80", "casos das ERs\n8 aceitas + 8 rejeitadas por ER"),
               ("15.000", "cadeias comparando\nAFNε × re.fullmatch"), ("0", "falhas ou\ndivergências")]
    for (n, rotulo), x, w in zip(numeros, (0.6, 3.3, 6.0, 10.2), (2.5, 2.5, 3.9, 2.5)):
        texto(s, x, 1.8, w, 1.1, n, tamanho=54, negrito=True, fonte=TITULO, cor_=AMBAR)
        texto(s, x, 2.9, w, 0.9, rotulo, tamanho=14, cor_="CADCFC")
    texto(s, 0.6, 4.25, 12.1, 0.4, "FALSOS RESULTADOS CONHECIDOS (a ER valida o formato, não a aritmética)",
          tamanho=13, negrito=True, cor_=AMBAR)
    texto(s, 0.6, 4.7, 12.1, 2.0, [
        [("ER-01  ", {"negrito": True, "fonte": MONO}), ("00.000.000/0000-00 é aceito: dígitos verificadores não são calculados.", {})],
        [("ER-03  ", {"negrito": True, "fonte": MONO}), ("DDD existente com número inativo é aceito.", {})],
        [("ER-04  ", {"negrito": True, "fonte": MONO}), ("29/02/2023 é aceito: a ER não distingue anos bissextos.", {})],
    ], tamanho=16, cor_=BRANCO, espaco_depois=8)
    rodape(s, "Detalhes: docs/RESULTADOS_TESTES.md · página Testes da aplicação · python -m pytest -v", claro=True)
    notas(s, "JOÃO: Cada caso é testado duas vezes, com a regex e com o AFNε. A comparação em 15.000 cadeias "
             "geradas por mutação mostra que autômato e código aceitam a mesma linguagem. Os falsos positivos "
             "são limitações documentadas, não erros.")


def slide_limitacoes(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, "Limitações e possíveis melhorias")
    for i, (cab, itens, tom) in enumerate([
        ("Limitações", [
            "Formato ≠ existência: DV do CNPJ, placa registrada e telefone ativo não são consultados.",
            "29/02 aceito em qualquer ano (ER-04).",
            "Sensível a maiúsculas e espaços, de propósito, para não alterar a linguagem.",
            "AFNε de Thompson são grandes (84 estados na ER-03).",
        ], VERMELHO),
        ("Melhorias", [
            "Validação semântica após a ER: DV módulo 11 e anos bissextos.",
            "Converter AFNε em AFD (subconjuntos) e minimizar.",
            "Normalização opcional com aviso ao usuário.",
            "Novas ERs: CEP, chave da NF-e, RNTRC; planilhas .xlsx.",
        ], VERDE),
    ]):
        x = 0.6 + i * 6.15
        caixa(s, x, 1.55, 5.95, 4.9, TINTA)
        texto(s, x + 0.3, 1.8, 5.4, 0.5, cab, tamanho=24, negrito=True, fonte=TITULO, cor_=tom)
        texto(s, x + 0.3, 2.6, 5.4, 3.7, itens, tamanho=18, marcadores=True, espaco_depois=14)
    notas(s, "JOÃO: Limitações são decisões conscientes: a ER descreve formato. As melhorias mostram como "
             "evoluir o trabalho, incluindo conteúdos da disciplina (AFD e minimização).")


def slide_equipe(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, BRANCO)
    titulo(s, "Contribuições da equipe")
    membros = [
        ("Yuri Aguiar", "Arquitetura, as 5 ERs, AFNε (Thompson), validador, CSV, interface web, testes "
                        "automatizados, documentação", "Thompson · ER-03 · demonstração"),
        ("Pedro Paulo", "Dados de exemplo, casos de teste das ER-01 e ER-02, revisão da interface",
         "ER-01 · ER-02"),
        ("João Rath", "Casos de teste das ER-04 e ER-05, revisão da notação formal e do relatório",
         "ER-04 · ER-05 · testes"),
    ]
    for i, (nome, fez, apresenta) in enumerate(membros):
        x = 0.6 + i * 4.1
        caixa(s, x, 1.6, 3.85, 3.3, TINTA)
        caixa(s, x + 0.3, 1.85, 0.8, 0.8, MARINHO, raio=0.5)
        iniciais = "".join(p[0] for p in nome.split())
        texto(s, x + 0.3, 1.85, 0.8, 0.8, iniciais, tamanho=20, negrito=True, cor_=AMBAR,
              alinhamento=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
        texto(s, x + 1.25, 1.85, 2.5, 0.8, nome, tamanho=20, negrito=True, fonte=TITULO, cor_=MARINHO,
              ancora=MSO_ANCHOR.MIDDLE)
        texto(s, x + 0.3, 2.85, 3.3, 1.3, fez, tamanho=14)
        texto(s, x + 0.3, 4.2, 3.3, 0.5, apresenta, tamanho=13, negrito=True, cor_=CINZA)
    caixa(s, 0.6, 5.25, 12.1, 1.25, "FFF6E0")
    texto(s, 0.85, 5.37, 11.6, 1.05, [
        [("Em conjunto: ", {"negrito": True}), ("definição do problema, escolha das ERs e revisão dos AFNε.", {})],
        [("Uso de IA: ", {"negrito": True}),
         ("Claude Code (Anthropic) apoiou a estrutura do projeto, o módulo do AFNε, os scripts de documentação e "
          "a redação inicial de README, relatório e slides. Todo o conteúdo foi revisado e é explicado pela equipe.", {})],
    ], tamanho=13, espaco_depois=6)
    notas(s, "YURI: Apresentar a contribuição de cada integrante e a "
             "declaração de uso de IA exigida pelo professor.")


def slide_final(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, MARINHO)
    texto(s, 0.8, 2.3, 11.7, 1.2, "Obrigado!", tamanho=60, negrito=True, fonte=TITULO, cor_=BRANCO)
    texto(s, 0.8, 3.5, 11.7, 0.7, "Perguntas?", tamanho=28, cor_=AMBAR)
    texto(s, 0.8, 5.2, 11.7, 1.0, ["Repositório: github.com/yuriaguiardev/valida-frete",
                                   "Yuri Aguiar · Pedro Paulo · João Rath — CC6NA"], tamanho=16, cor_="CADCFC",
          espaco_depois=6)
    notas(s, "TODOS: Abrir para perguntas. Deixar a aplicação aberta para o professor sugerir novas entradas.")


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_capa(prs)
    slide_problema(prs)
    slide_solucao(prs)
    slide_fonte_unica(prs)
    slide_thompson(prs)

    oradores = {"cnpj": "PEDRO", "placa": "PEDRO", "telefone": "YURI", "data": "JOÃO", "valor": "JOÃO"}
    for er in EXPRESSOES:
        slide_ficha(prs, er, oradores[er.chave])
        if er.chave == "telefone":
            slide_afn(prs, er, oradores[er.chave], com_testes=False)
            slide_testes_tel(prs, er, oradores[er.chave])
        else:
            slide_afn(prs, er, oradores[er.chave], partes=2 if er.chave == "cnpj" else 1)

    slide_entradas(prs)
    slide_demo(prs)
    slide_numeros(prs)
    slide_limitacoes(prs)
    slide_equipe(prs)
    slide_final(prs)

    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    prs.save(SAIDA)
    print(f"{SAIDA.relative_to(RAIZ)}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
